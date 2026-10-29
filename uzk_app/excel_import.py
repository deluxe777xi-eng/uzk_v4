# -*- coding: utf-8 -*-
"""Импорт данных контроля из Excel (openpyxl для .xlsx, xlrd для .xls)."""
import os

class ExcelParseError(Exception):
    """Файл не читается или не соответствует ожидаемой структуре."""

MAX_SCREENSHOT_SIZE = 10 * 1024 * 1024

def read_rows(file_path):
    if not os.path.exists(file_path):
        raise ExcelParseError(f"Файл не найден: {file_path}")
    try:
        if file_path.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
            return [list(row) for row in ws.iter_rows(values_only=True)]
        else:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            sh = wb.sheet_by_index(0)
            return [sh.row_values(i) for i in range(sh.nrows)]
    except ImportError as e:
        raise ExcelParseError("Библиотеки openpyxl или xlrd не установлены.") from e
    except ExcelParseError:
        raise
    except Exception as e:
        raise ExcelParseError(f"Не удалось прочитать файл: {e}") from e

def safe_str(val):
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        if isinstance(val, float) and val.is_integer():
            return str(int(val))
        return str(val).strip()
    return str(val).strip()

def split_fio_values(value):
    fio_values = []
    for fio in str(value or "").replace("\r", "\n").split("\n"):
        for item in fio.replace(";", ",").split(","):
            item = item.strip()
            if item and item not in fio_values:
                fio_values.append(item)
    return fio_values

def parse_excel(file_path):
    rows = read_rows(file_path)
    if not rows:
        return None

    header_row_idx = None
    for i, row in enumerate(rows):
        if row and safe_str(row[0]) == "Изделие" and safe_str(row[1]) == "Деталь":
            header_row_idx = i
            break

    if header_row_idx is None or header_row_idx + 1 >= len(rows):
        raise ExcelParseError("Не найдена строка с заголовком 'Изделие/Деталь'")

    data_row = rows[header_row_idx + 1]
    general = {
        'izdeliye': safe_str(data_row[0]) if len(data_row) > 0 else "",
        'detal': safe_str(data_row[1]) if len(data_row) > 1 else "",
        'object_kontr': safe_str(data_row[2]) if len(data_row) > 2 else "",
        'thickness': safe_str(data_row[3]) if len(data_row) > 3 else "",
    }
    total_plots_str = safe_str(data_row[5]) if len(data_row) > 5 else "0"
    method = safe_str(data_row[6]) if len(data_row) > 6 else ""
    try:
        total_plots = int(float(total_plots_str.replace(',', '.')))
    except Exception:
        total_plots = 0
    general['total_plots'] = total_plots
    general['method'] = method

    channel_start = None
    for i, row in enumerate(rows):
        if row and safe_str(row[0]).lower() == "канал":
            channel_start = i
            break
    if channel_start is None:
        raise ExcelParseError("Не найдена строка с заголовком 'Канал'")

    def parse_equip(val):
        if "№" in val:
            parts = val.split("№", 1)
            return parts[0].strip(), parts[1].strip()
        return val, ""

    channel_map = {}
    equipment_to_add = []
    problems = []
    idx = channel_start + 1
    while idx < len(rows):
        row = rows[idx]
        if not row or not safe_str(row[0]):
            break
        ch = safe_str(row[0])
        if not ch:
            idx += 1
            continue

        probe_type, probe_sn = parse_equip(safe_str(row[1]) if len(row) > 1 else "")
        prism_type, prism_sn = parse_equip(safe_str(row[2]) if len(row) > 2 else "")

        if probe_type:
            channel_map.setdefault(ch, []).append({
                'probe_type': probe_type, 'probe_sn': probe_sn,
                'prism_type': prism_type, 'prism_sn': prism_sn,
            })
            for cat, t, s in (('probe', probe_type, probe_sn), ('prism', prism_type, prism_sn)):
                if t:
                    equipment_to_add.append((cat, t, s))
        else:
            problems.append((idx + 1, "Не указан тип преобразователя; строка канала пропущена."))
        idx += 1

    defect_start = None
    for i, row in enumerate(rows):
        if row and safe_str(row[0]).startswith("№ дефекта"):
            defect_start = i
            break
    if defect_start is None:
        raise ExcelParseError("Не найдена строка с заголовком '№ дефекта'")

    defect_header = rows[defect_start]
    fio_column_idx = next(
        (i for i, value in enumerate(defect_header) if "фио" in safe_str(value).lower()), 15)
    len_x_column_idx = next(
        (i for i, value in enumerate(defect_header)
         if "протяженность" in safe_str(value).lower()
         and ("x" in safe_str(value).lower() or "х" in safe_str(value).lower())), 18)

    defects_data = []
    idx = defect_start + 1
    while idx < len(rows):
        row = rows[idx]
        if not row or not safe_str(row[1]):
            break
        plot = safe_str(row[1])
        if not plot:
            idx += 1
            continue

        adef = safe_str(row[10]) if len(row) > 10 else ""
        sdef = safe_str(row[11]) if len(row) > 11 else ""
        defects_data.append({
            '_source_row': idx + 1,
            'channel': safe_str(row[3]) if len(row) > 3 else "",
            'plot_num': plot,
            'defect_type': safe_str(row[2]) if len(row) > 2 else "",
            'z_start_end': safe_str(row[4]) if len(row) > 4 else "",
            'x_start_end': safe_str(row[5]) if len(row) > 5 else "",
            'y_start_end': safe_str(row[6]) if len(row) > 6 else "",
            'height': safe_str(row[7]) if len(row) > 7 else "",
            'len_y': safe_str(row[8]) if len(row) > 8 else "",
            'quality': safe_str(row[9]) if len(row) > 9 else "",
            'len_x': safe_str(row[len_x_column_idx]) if len(row) > len_x_column_idx else "",
            'fio': ", ".join(split_fio_values(row[fio_column_idx]) if len(row) > fio_column_idx else []),
            'amp_s_def': f"{adef} / {sdef}",
            'no_defect': "",
            'defect_object': general['object_kontr'],
        })
        idx += 1

    if not channel_map:
        raise ExcelParseError("Не найдено ни одного канала/преобразователя")

    for defect in defects_data:
        if defect["channel"] not in channel_map:
            problems.append((
                defect["_source_row"],
                f"Канал «{defect['channel']}» не найден в таблице оборудования; дефект не будет импортирован.",
            ))

    return {
        'general': general,
        'channel_map': channel_map,
        'equipment_to_add': equipment_to_add,
        'defects_data': defects_data,
        'problems': problems,
        'method': method,
    }

def _occupied_plots(raw_plot):
    occupied = set()
    for part in raw_plot.split(','):
        part = part.strip()
        if not part:
            continue
        part = part.replace(',', '.').replace('÷', '-')
        if '-' in part:
            try:
                start_str, end_str = part.split('-')
                start = int(float(start_str.replace(',', '.')))
                end = int(float(end_str.replace(',', '.')))
                for p in range(start, end + 1):
                    occupied.add(p)
            except Exception:
                pass
        else:
            try:
                occupied.add(int(float(part.replace(',', '.'))))
            except Exception:
                pass
    return occupied

def _free_plots_range_string(free_plots):
    ranges = []
    start = free_plots[0]
    end = start
    for i in range(1, len(free_plots)):
        if free_plots[i] == end + 1:
            end = free_plots[i]
        else:
            ranges.append(str(start) if start == end else f"{start}÷{end}")
            start = free_plots[i]
            end = start
    ranges.append(str(start) if start == end else f"{start}÷{end}")
    return ", ".join(ranges)

def build_measurements(parsed, current_object_value="", existing_measurements=None):
    general = parsed['general']
    channel_map = parsed['channel_map']
    defects_data = parsed['defects_data']
    object_kontr = general['object_kontr']
    total_plots = general['total_plots']

    grouped = {}
    for d in defects_data:
        grouped.setdefault(d['channel'], []).append(d)
    for ch in channel_map:
        grouped.setdefault(ch, [])

    imported_fios = []
    for defect in defects_data:
        for fio in split_fio_values(defect.get('fio')):
            if fio not in imported_fios:
                imported_fios.append(fio)
    imported_fio = ", ".join(imported_fios)

    new_measurements = []
    for ch, defects in grouped.items():
        if ch not in channel_map:
            continue
        equip_list = channel_map[ch]
        equip1 = equip_list[0] if len(equip_list) > 0 else None
        equip2 = equip_list[1] if len(equip_list) > 1 else None

        m_data = {
            "tip_preobr": equip1['probe_type'] if equip1 else "",
            "tip_preobr_sn": equip1['probe_sn'] if equip1 else "",
            "tip_prism": equip1['prism_type'] if equip1 else "",
            "tip_prism_sn": equip1['prism_sn'] if equip1 else "",
            "tip_preobr2": equip2['probe_type'] if equip2 else "",
            "tip_preobr2_sn": equip2['probe_sn'] if equip2 else "",
            "tip_prism2": equip2['prism_type'] if equip2 else "",
            "tip_prism2_sn": equip2['prism_sn'] if equip2 else "",
            "kanal": ch, "tip_skanirov": "", "time_delay": "", "angle": "",
            "focus": "", "dist_points": "", "preamp": "", "averaging": "",
            "offset_y": "", "scan_range_x": "", "sample_ref": "", "setup_method": "",
            "we_level": "", "gain": "",
            "defects": [],
            "_import_object": object_kontr,
        }

        occupied_plots = set()
        for d in defects:
            occupied_plots |= _occupied_plots(d['plot_num'])

        if total_plots > 0:
            all_plots = set(range(1, total_plots + 1))
            free_plots = sorted(all_plots - occupied_plots)
            if free_plots:
                defects.append({
                    'plot_num': _free_plots_range_string(free_plots),
                    'no_defect': "Да",
                    'quality': "уд.",
                    'fio': imported_fio,
                    'defect_object': object_kontr,
                })

        for defect in defects:
            defect.pop('_source_row', None)
        m_data['defects'] = defects
        new_measurements.append(m_data)

    ordered_objects = []
    current_field_value = (current_object_value or "").strip()
    if current_field_value:
        for part in current_field_value.split(","):
            part = part.strip()
            if part and part not in ordered_objects:
                ordered_objects.append(part)

    for m in (existing_measurements or []):
        for d in m.get('defects', []):
            if not isinstance(d, dict):
                continue
            dobj = (d.get('defect_object') or '').strip()
            for part in dobj.split(","):
                p = part.strip()
                if p and p not in ordered_objects:
                    ordered_objects.append(p)

    for obj in [o.strip() for o in (object_kontr or "").split(",") if o.strip()]:
        if obj not in ordered_objects:
            ordered_objects.append(obj)

    # ИСПРАВЛЕНИЕ: Удалены строки, которые принудительно перезаписывали d['defect_object']
    
    return new_measurements, ordered_objects